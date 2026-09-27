#!/usr/bin/env python3
"""The de-risking ledger (SPEC.md §5.13): the package's stand-in for
`cargo xtask derisk` over the crate adcs-derisk.

    python3 tools/derisk.py check                      the ledger keeps its rules
    python3 tools/derisk.py rollup [--quarter Q3-26] [--json]
                                                       the values of the 18 risk rows, and the 3 conclusions
    python3 tools/derisk.py narrative --quarter Q3-26 --out <dir>
                                                       narrative_<Q>.html, .xlsx and .csv
    python3 tools/derisk.py template --out <file.xlsx> the company's blank narrative template
    python3 tools/derisk.py table [--check] [--file <path>]
                                                       the register as a table: into SPEC.md §21 in this
                                                       package; into docs/RISKS.md in the repository
    python3 tools/derisk.py record --area <a> --about <x,y> --believed <text> --status broke|held|untested
                          [--tested <text>] [--would-test <text>] --now-know <text> --plan-change <text>
                          [--previous-issue <text>] [--benefit <text>] [--cost-k <n>] [--quarter Qn-YY]
                          [--move R-nn:<from>:<to>]... [--risk R-nn]... [--open "<title>|<area>|<level>|<closing test>|<owner>"]...
                                                       a belief record for the developer team's own change,
                                                       written with its moves; refused if the ledger would
                                                       then fail its check
    python3 tools/derisk.py codes                      every rule, by code
    python3 tools/derisk.py selftest                   each deliberate mistake is refused by its rule

The ledger is derisk/risks.toml (the register) and derisk/beliefs/*.toml (one
belief record each). A person's prose for a quarter is derisk/narratives/<Q>.md.
Beliefs arrive through intake (a node form's De-risking section) or, for the
developer team's own changes, `record` (`cargo xtask derisk record` in the
repository). No other command writes the ledger.
"""

import copy
import csv
import html
import io
import json
import os
import re
import shutil
import sys
import tempfile

try:
    import tomllib
except ImportError:  # pragma: no cover
    raise SystemExit("the package tools need Python 3.11+ (tomllib)")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

AREAS = ["node", "input", "output", "model", "math", "algorithm", "visualisation"]
STATUSES = ["broke", "held", "untested"]
QUARTER = re.compile(r"^Q([1-4])-(\d\d)$")
RISK_ID = re.compile(r"^R-\d{2,}$")
BELIEF_ID = re.compile(r"^(B-\d{3,}|REQ-[0-9]{4}-[0-9]{2}-[0-9]{2}-[a-z0-9_]+-[0-9a-f]{6})$")
TEAMS = ["actuators", "avionics", "environment", "facility", "gnc", "programme", "quality", "sales", "sensing", "systems", "verification"]

# The layer-1 risk rows, by tree id (plan/tree.json), and what each counts.
ROWS = [
    ("rk1_0", "Risks open", "open"),
    ("rk1_1", "Risks open at level L4 or L5", "high"),
    ("rk1_2", "Risks closed", "closed"),
    ("rk1_3", "Risks opened this quarter", "opened_q"),
    ("rk1_4", "Risks closed this quarter", "closed_q"),
    ("rk2_0", "Beliefs a test broke", "broke"),
    ("rk2_1", "Beliefs a test held", "held"),
    ("rk2_2", "Beliefs not yet tested", "untested"),
    ("rk2_3", "Beliefs recorded", "beliefs"),
    ("rk2_4", "Node versions released", "versions"),
    ("rk2_5", "Cost of the tests this quarter (USD)", "cost_usd"),
] + [("rk3_%d" % i, "Highest open risk: %s" % a, "area:" + a) for i, a in enumerate(AREAS)]
CONCLUSION = [
    ("rk4_0", "Highest open risk level", "risk::highest_level over the seven areas"),
    ("rk4_1", "Net risks closed this quarter", "risk::net_closed(closed this quarter, opened this quarter)"),
    ("rk4_2", "Share of beliefs tested", "risk::share_tested(untested, recorded)"),
]

CODES = {
    "L01": "the register is adcs-risk-register/1; every risk id is R-nn and unique; every belief is adcs-belief/1 with a unique id, B-nnn or a request id, named as its file",
    "L02": "a risk's area is one of the seven; its owner is a team, never a person; its level is a whole number 0 to 5",
    "L03": "a risk's history starts from 0, each move starts where the last ended, the last ends at its level, and opened and closed are the quarters of the first move and of the move to 0",
    "L04": "every move in the register names the belief that made it (only a risk's first entry, made at planning, may name none), is written in that belief, and every move a belief makes is written in the register",
    "L05": "a level goes down only by a tested belief: its status is held or broke and it says what was tested",
    "L06": "a belief's area is one of the seven; its status is broke, held or untested; it says what was believed, what we now know and what changed in the plan",
    "L07": "a held or broke belief says what tested it; an untested belief says the test that would settle it",
    "L08": "quarters are Qn-YY; a cost is blank or a number, zero or more, in thousands of dollars",
    "L09": "every risk a belief names exists; every open risk says the test that would close it",
}


def load_toml(p):
    with open(p, "rb") as f:
        return tomllib.load(f)


def qkey(q):
    m = QUARTER.match(q or "")
    return (int(m.group(2)), int(m.group(1))) if m else (0, 0)


def load(root=None):
    root = root or ROOT
    reg = load_toml(os.path.join(root, "derisk", "risks.toml"))
    beliefs = {}
    bdir = os.path.join(root, "derisk", "beliefs")
    for name in sorted(os.listdir(bdir)) if os.path.isdir(bdir) else []:
        if name.endswith(".toml"):
            b = load_toml(os.path.join(bdir, name))
            b["_file"] = name
            beliefs[b.get("id", name[:-5])] = b
    prose = {}
    ndir = os.path.join(root, "derisk", "narratives")
    for name in sorted(os.listdir(ndir)) if os.path.isdir(ndir) else []:
        if name.endswith(".md") and QUARTER.match(name[:-3]):
            prose[name[:-3]] = open(os.path.join(ndir, name), encoding="utf-8").read()
    return reg, beliefs, prose


def moves_of(b):
    return [m for m in b.get("move", [])]


# ------------------------------------------------------------------ check
def check(reg, beliefs):
    F = []
    add = lambda code, where, text: F.append((code, where, text))
    if reg.get("schema") != "adcs-risk-register/1":
        add("L01", "risks.toml", "schema is %r" % reg.get("schema"))
    risks = {}
    for r in reg.get("risk", []):
        rid = r.get("id", "")
        if not RISK_ID.match(rid):
            add("L01", rid or "?", "risk id %r is not R-nn" % rid)
        if rid in risks:
            add("L01", rid, "risk id used twice")
        risks[rid] = r
    for bid, b in beliefs.items():
        if b.get("schema") != "adcs-belief/1":
            add("L01", bid, "schema is %r" % b.get("schema"))
        if not BELIEF_ID.match(bid):
            add("L01", bid, "belief id %r is neither B-nnn nor a request id" % bid)
        if b.get("_file") and b["_file"] != bid + ".toml":
            add("L01", bid, "belief %s is in the file %s" % (bid, b["_file"]))

    for rid, r in risks.items():
        if r.get("area") not in AREAS:
            add("L02", rid, "area %r is not one of %s" % (r.get("area"), ", ".join(AREAS)))
        if r.get("owner") not in TEAMS:
            add("L02", rid, "owner %r is not a team" % r.get("owner"))
        lvl = r.get("level")
        if not isinstance(lvl, int) or isinstance(lvl, bool) or not 0 <= lvl <= 5:
            add("L02", rid, "level %r is not a whole number 0 to 5" % (lvl,))
            continue
        hist = r.get("history", [])
        if not hist:
            add("L03", rid, "no history: a risk is entered by a move from 0")
            continue
        if hist[0].get("from") != 0:
            add("L03", rid, "the first move starts at %r, not 0" % hist[0].get("from"))
        for a, b in zip(hist, hist[1:]):
            if b.get("from") != a.get("to"):
                add("L03", rid, "a move starts at %r where the last ended at %r" % (b.get("from"), a.get("to")))
            if qkey(b.get("quarter")) < qkey(a.get("quarter")):
                add("L03", rid, "the history is not in quarter order")
        if hist[-1].get("to") != lvl:
            add("L03", rid, "the last move ends at %r, the level is %r" % (hist[-1].get("to"), lvl))
        if r.get("opened") != hist[0].get("quarter"):
            add("L03", rid, "opened %r, first move in %r" % (r.get("opened"), hist[0].get("quarter")))
        closes = [h for h in hist if h.get("to") == 0]
        want_closed = closes[-1].get("quarter") if (lvl == 0 and closes) else ""
        if (r.get("closed") or "") != want_closed:
            add("L03", rid, "closed %r, but the history says %r" % (r.get("closed"), want_closed))
        for i, h in enumerate(hist):
            if not QUARTER.match(h.get("quarter", "")):
                add("L08", rid, "quarter %r is not Qn-YY" % h.get("quarter"))
            bid = h.get("belief", "")
            if not bid and not (i == 0 and "planning" in h.get("note", "")):
                add("L04", rid, "the move %s→%s in %s names no belief; only a risk's first entry, made at planning, may" % (h.get("from"), h.get("to"), h.get("quarter")))
            if bid:
                b = beliefs.get(bid)
                if not b:
                    add("L04", rid, "a move names belief %s, which is not in derisk/beliefs/" % bid)
                    continue
                if not any(m.get("risk") == rid and m.get("from") == h.get("from") and m.get("to") == h.get("to") for m in moves_of(b)) \
                        or b.get("quarter") != h.get("quarter"):
                    add("L04", rid, "the move %s→%s in %s is not written in belief %s" % (h.get("from"), h.get("to"), h.get("quarter"), bid))
            if isinstance(h.get("to"), int) and isinstance(h.get("from"), int) and h["to"] < h["from"]:
                b = beliefs.get(bid) if bid else None
                if not b or b.get("status") not in ("held", "broke") or not str(b.get("tested", "")).strip():
                    add("L05", rid, "lowered %s→%s in %s without a tested belief" % (h["from"], h["to"], h.get("quarter")))
        if lvl > 0 and not str(r.get("closing_test", "")).strip():
            add("L09", rid, "an open risk must say the test that would close it")

    for bid, b in beliefs.items():
        if b.get("area") not in AREAS:
            add("L06", bid, "area %r is not one of %s" % (b.get("area"), ", ".join(AREAS)))
        if b.get("status") not in STATUSES:
            add("L06", bid, "status %r is not broke, held or untested" % b.get("status"))
        for f in ("believed", "now_know", "plan_change"):
            if not str(b.get(f, "")).strip():
                add("L06", bid, "%s is empty" % f)
        if b.get("status") in ("held", "broke") and not str(b.get("tested", "")).strip():
            add("L07", bid, "a %s belief must say what tested it" % b.get("status"))
        if b.get("status") == "untested" and not str(b.get("would_test", "")).strip():
            add("L07", bid, "an untested belief must say the test that would settle it")
        if not QUARTER.match(b.get("quarter", "")):
            add("L08", bid, "quarter %r is not Qn-YY" % b.get("quarter"))
        c = b.get("cost_k", "")
        if c != "" and (isinstance(c, bool) or not isinstance(c, (int, float)) or c < 0):
            add("L08", bid, "cost %r is neither blank nor a number, zero or more" % (c,))
        for rid in list(b.get("risks", [])) + [m.get("risk") for m in moves_of(b)]:
            if rid not in risks:
                add("L09", bid, "names risk %s, which is not in the register" % rid)
        for m in moves_of(b):
            r = risks.get(m.get("risk"))
            if r and not any(h.get("belief") == bid and h.get("from") == m.get("from") and h.get("to") == m.get("to")
                             and h.get("quarter") == b.get("quarter") for h in r.get("history", [])):
                add("L04", bid, "its move of %s %s→%s is not written in the register" % (m.get("risk"), m.get("from"), m.get("to")))
    return F


# ------------------------------------------------------------------ rollup
def latest_quarter(reg, beliefs):
    qs = [h.get("quarter") for r in reg.get("risk", []) for h in r.get("history", [])] + [b.get("quarter") for b in beliefs.values()]
    qs = [q for q in qs if QUARTER.match(q or "")]
    return max(qs, key=qkey) if qs else ""


def count_versions(root):
    """Node versions released: versions.toml entries with a release, in the
    repository's node folders. The package has no node folders, so none."""
    n = 0
    crates = os.path.join(root, "crates")
    for dirpath, _, files in os.walk(crates) if os.path.isdir(crates) else []:
        if "versions.toml" in files:
            n += sum(1 for v in load_toml(os.path.join(dirpath, "versions.toml")).get("version", []) if v.get("release"))
    return n


def rollup(reg, beliefs, quarter=None, root=None):
    quarter = quarter or latest_quarter(reg, beliefs)
    risks = reg.get("risk", [])
    v = {
        "open": sum(1 for r in risks if r.get("level", 0) > 0),
        "high": sum(1 for r in risks if r.get("level", 0) >= 4),
        "closed": sum(1 for r in risks if r.get("level", 0) == 0),
        "opened_q": sum(1 for r in risks if r.get("opened") == quarter),
        "closed_q": sum(1 for r in risks if r.get("closed") == quarter),
        "broke": sum(1 for b in beliefs.values() if b.get("status") == "broke"),
        "held": sum(1 for b in beliefs.values() if b.get("status") == "held"),
        "untested": sum(1 for b in beliefs.values() if b.get("status") == "untested"),
        "beliefs": len(beliefs),
        "versions": count_versions(root or ROOT),
        "cost_usd": 1000.0 * sum(float(b.get("cost_k") or 0) for b in beliefs.values() if b.get("quarter") == quarter),
    }
    for a in AREAS:
        v["area:" + a] = max([r.get("level", 0) for r in risks if r.get("area") == a and r.get("level", 0) > 0] or [0])
    rows = {tid: v[key] for tid, _, key in ROWS}
    concl = {
        "rk4_0": max(v["area:" + a] for a in AREAS),
        "rk4_1": v["closed_q"] - v["opened_q"],
        "rk4_2": (1.0 - v["untested"] / v["beliefs"]) if v["beliefs"] else None,
    }
    return quarter, rows, concl, v


# ------------------------------------------------------------------ narrative
COLUMNS = ["Quarter", "What we believed", "What we tested", "What we now know", "What it cost ($k)",
           "What changed in the plan", "Risks opened / closed"]
TITLE = "Quarterly De-risking Narrative"
SUBTITLE = "One page of prose per quarter. Bad news appears here first and between meetings — never sprung at the table."


def moves_text(b, risks):
    out = []
    for m in moves_of(b):
        rid, a, z = m.get("risk"), m.get("from"), m.get("to")
        if a == 0:
            out.append("Opened %s (L%d)" % (rid, z))
        elif z == 0:
            out.append("Closed %s" % rid)
        elif z < a:
            out.append("%s reduced L%d->L%d" % (rid, a, z))
        else:
            out.append("%s raised L%d->L%d" % (rid, a, z))
    return "; ".join(out) or "none"


def table_rows(beliefs, quarter, risks):
    rows = []
    for bid in sorted(beliefs):
        b = beliefs[bid]
        if b.get("quarter") != quarter:
            continue
        tested = str(b.get("tested", "")).strip() or "Not yet tested. The test that would settle it: " + str(b.get("would_test", "")).strip()
        cost = b.get("cost_k", "")
        rows.append([quarter, b.get("believed", ""), tested, b.get("now_know", ""), "" if cost == "" else cost,
                     b.get("plan_change", ""), moves_text(b, risks), bid, b.get("status", ""), b.get("area", "")])
    return rows


def write_xlsx(path, rows, blank=False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "De-risking Narrative"
    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:G1")
    ws["A2"] = SUBTITLE
    ws["A2"].font = Font(size=9)
    ws.merge_cells("A2:G2")
    head = PatternFill("solid", fgColor="FF0F2B46")
    for i, c in enumerate(COLUMNS, 1):
        cell = ws.cell(row=4, column=i, value=c)
        cell.font = Font(bold=True, size=10, color="FFFFFFFF")
        cell.fill = head
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r, row in enumerate([] if blank else rows, 5):
        for i, val in enumerate(row[:7], 1):
            cell = ws.cell(row=r, column=i, value=val)
            cell.font = Font(size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFG", (10, 40, 40, 44, 14, 40, 22)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wb.properties.creator = "tools/derisk.py"
    wb.save(path)


def write_csv(path, rows):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(COLUMNS + ["Belief", "Status", "Area"])
        for row in rows:
            w.writerow(row)


LEVEL_WORD = {0: "none open", 1: "negligible", 2: "minor", 3: "significant", 4: "major", 5: "critical"}


def narrative_html(reg, beliefs, prose, quarter):
    from explain_kit import page_head
    css_kit, js_kit = page_head()
    e = lambda x: html.escape(str(x), quote=True)
    risks = {r["id"]: r for r in reg.get("risk", [])}
    q, rows, concl, v = rollup(reg, beliefs, quarter)
    area_label = {a["id"]: a["label"] for a in reg.get("area", [])}
    trows = table_rows(beliefs, quarter, risks)
    untested = [b for b in beliefs.values() if b.get("status") == "untested"]
    high = sorted([r for r in risks.values() if r.get("level", 0) >= 4], key=lambda r: (-r["level"], r["id"]))
    share = concl["rk4_2"]
    share_txt = "no beliefs recorded" if share is None else "%d of %d beliefs tested" % (v["beliefs"] - v["untested"], v["beliefs"])
    answer = ("In %s the highest open risk anywhere in the platform is <b>L%d, %s</b>; the quarter closed %d risk%s and opened %d "
              "(net %+d); and %s. Every level is still a planning proposal until decision D26 confirms it."
              % (e(q), concl["rk4_0"], LEVEL_WORD[concl["rk4_0"]], v["closed_q"], "" if v["closed_q"] == 1 else "s",
                 v["opened_q"], concl["rk4_1"], share_txt))
    L = []
    A = L.append
    A("<!doctype html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">")
    A("<title>De-risking narrative %s</title>" % e(q))
    A("<!-- Written by tools/derisk.py narrative from derisk/ (SPEC.md §5.13). Follows adcs-explain/1 (§5.12). -->")
    A("<style>")
    A(""":root { color-scheme: light; --page:#f9f9f7; --surface:#fcfcfb; --raised:#fff; --ink:#0b0b0b; --ink-2:#52514e; --muted:#898781;
  --line:#e1e0d9; --line-2:#c3c2b7; --accent:#2a78d6; --accent-ink:#1c5cab; --accent-soft:#e6f0fc; --good:#0ca30c; --good-ink:#006300;
  --bad:#d03b3b; --warn-ink:#7a5300; --bar:#2a78d6; }
@media (prefers-color-scheme: dark) { :root:not([data-theme="light"]) { color-scheme: dark; --page:#0d0d0d; --surface:#1a1a19; --raised:#222220;
  --ink:#fff; --ink-2:#c3c2b7; --muted:#898781; --line:#2c2c2a; --line-2:#383835; --accent:#3987e5; --accent-ink:#86b6ef; --accent-soft:#16263a;
  --good:#0ca30c; --good-ink:#4cc44c; --bad:#e66767; --warn-ink:#f0c060; --bar:#3987e5; } }
:root[data-theme="dark"] { color-scheme: dark; --page:#0d0d0d; --surface:#1a1a19; --raised:#222220; --ink:#fff; --ink-2:#c3c2b7; --muted:#898781;
  --line:#2c2c2a; --line-2:#383835; --accent:#3987e5; --accent-ink:#86b6ef; --accent-soft:#16263a; --good:#0ca30c; --good-ink:#4cc44c;
  --bad:#e66767; --warn-ink:#f0c060; --bar:#3987e5; }
* { box-sizing: border-box; }
body { margin: 0; background: var(--page); color: var(--ink); font: 15px/1.55 system-ui, -apple-system, "Segoe UI", Roboto, sans-serif; }
header.top { position: sticky; top: 0; z-index: 5; background: var(--surface); border-bottom: 1px solid var(--line); }
.top .bar { max-width: 1100px; margin: 0 auto; padding: 10px 16px; display: flex; gap: 10px; align-items: center; flex-wrap: wrap; }
.top h1 { font-size: 16px; margin: 0; flex: 1 1 260px; }
main { max-width: 1100px; margin: 0 auto; padding: 16px; }
section.card { background: var(--surface); border: 1px solid var(--line); border-radius: 12px; padding: 16px 18px; margin: 0 0 16px; }
section.card > h2 { font-size: 17px; margin: 0 0 6px; }
.lead { color: var(--ink-2); margin: 0 0 10px; }
.tiles { display: grid; grid-template-columns: repeat(auto-fit, minmax(200px, 1fr)); gap: 10px; }
.tile { background: var(--raised); border: 1px solid var(--line); border-radius: 10px; padding: 10px 12px; }
.tile .k { font-size: 12px; color: var(--ink-2); }
.tile .v { font-size: 26px; font-weight: 600; font-variant-numeric: tabular-nums; }
.tile .s { font-size: 12px; color: var(--muted); }
.areas { display: grid; grid-template-columns: minmax(120px, 180px) minmax(0, 1fr) 150px; gap: 6px 10px; align-items: center; font-size: 13px; }
.areas .track { position: relative; height: 16px; border-radius: 4px; background: repeating-linear-gradient(90deg, var(--line) 0 1px, transparent 1px 20%); border: 1px solid var(--line); }
.areas .fill { position: absolute; left: 0; top: 0; bottom: 0; border-radius: 0 4px 4px 0; background: var(--bar); }
.axis { display: grid; grid-template-columns: minmax(120px, 180px) minmax(0, 1fr) 150px; gap: 10px; font-size: 11px; color: var(--muted); }
.axis .ticks { display: flex; justify-content: space-between; }
table { border-collapse: collapse; width: 100%; font-size: 13px; }
th, td { text-align: left; padding: 7px; border-bottom: 1px solid var(--line); vertical-align: top; }
th { background: #0f2b46; color: #fff; font-size: 12px; }
.scroll { overflow-x: auto; }
details { border: 1px solid var(--line); border-radius: 10px; padding: 10px 12px; background: var(--raised); margin: 8px 0; }
summary { cursor: pointer; font-weight: 600; }
textarea { width: 100%; min-height: 70px; font: inherit; padding: 8px; border: 1px solid var(--line-2); border-radius: 8px; background: var(--surface); color: var(--ink); }
button { font: inherit; font-size: 13px; padding: 6px 11px; border-radius: 7px; border: 1px solid var(--line-2); background: var(--raised); color: var(--ink); cursor: pointer; }
.prose { font-size: 15px; }
.missing { color: var(--ink-2); font-style: italic; }
.tiny { color: var(--muted); font-size: 12px; }
@media (max-width: 640px) { .areas, .axis { grid-template-columns: 1fr; } .axis { display: none; } }
@media print { header.top { position: static; } }""")
    A(css_kit)
    A("</style>\n<script>\n%s\n</script>\n</head>" % js_kit)
    A("<body>")
    A("<header class=\"top\"><div class=\"bar\"><h1>De-risking narrative · %s</h1><div data-depth-host></div>"
      "<button type=\"button\" onclick=\"(function(r){var d=r.getAttribute('data-theme')==='dark'||(!r.getAttribute('data-theme')&&matchMedia('(prefers-color-scheme: dark)').matches);r.setAttribute('data-theme',d?'light':'dark');})(document.documentElement)\" aria-label=\"Switch light or dark\">◐</button></div></header>" % e(q))
    A("<main>")
    A("<nav class=\"xk-crumbs\" data-explain=\"zoom\" aria-label=\"Where you are\"><span>ADCS platform</span><span>Layer 1 — the company</span><span>Risk management</span><span>Narrative %s</span></nav>" % e(q))
    A("<p class=\"xk-answer\" data-explain=\"answer\">%s</p>" % answer)

    # overview: the three conclusions
    A("<section class=\"card\" data-kind=\"explanation\" data-explain=\"overview\"><h2>The conclusion</h2>")
    A("<p class=\"lead\">Three numbers the risk branch computes for the whole platform (the Conclusion group, §5.13). Read them together: the first cannot move until the last risk at its level closes, so the other two show progress beneath it.</p>")
    A("<div class=\"tiles\">")
    A("<div class=\"tile\"><div class=\"k\">Highest open risk level <span class=\"xk-claim\" data-claim=\"derived\" title=\"Worked here from the register\">derived</span></div><div class=\"v\">L%d</div><div class=\"s\">%s · risk::highest_level</div></div>" % (concl["rk4_0"], LEVEL_WORD[concl["rk4_0"]]))
    A("<div class=\"tile\"><div class=\"k\">Net risks closed this quarter <span class=\"xk-claim\" data-claim=\"derived\">derived</span></div><div class=\"v\">%+d</div><div class=\"s\">%d closed, %d opened · risk::net_closed</div></div>" % (concl["rk4_1"], v["closed_q"], v["opened_q"]))
    A("<div class=\"tile\"><div class=\"k\">Share of beliefs tested <span class=\"xk-claim\" data-claim=\"derived\">derived</span></div><div class=\"v\">%s</div><div class=\"s\">%s · risk::share_tested</div></div>"
      % ("—" if share is None else "%d%%" % round(100 * share), share_txt))
    A("</div>")
    A("<div class=\"xk-box\" data-show=\"learn\" data-explain=\"predict\"><div class=\"xk-h\">Predict first</div>Before you read the next card: which of the seven areas do you expect to carry the highest open level, and why? Write it down, then look.</div>")
    A("</section>")

    # the prose
    A("<section class=\"card\" data-kind=\"explanation\"><h2>This quarter, in prose</h2>")
    if prose.get(q):
        paras = [p.strip() for p in prose[q].split("\n\n") if p.strip() and not p.strip().startswith("#")]
        A("<div class=\"prose\">" + "".join("<p>%s</p>" % e(p) for p in paras) + "</div>")
    else:
        A("<p class=\"missing\">No prose has been written for %s. The quality team writes it in <code>derisk/narratives/%s.md</code> (decision D27). This page does not write it for them: the numbers and the table below are counted, the judgement is a person's.</p>" % (e(q), e(q)))
    A("</section>")

    # by area, same axis
    A("<section class=\"card\" data-kind=\"reference\"><h2>Open risk by area</h2>")
    A("<p class=\"lead\">The highest open level in each area, all on the same scale from 0 (nothing open) to 5 (critical). These are the seven <i>Open risk by area</i> rows.</p>")
    A("<div class=\"axis\" aria-hidden=\"true\"><span></span><div class=\"ticks\"><span>0</span><span>1</span><span>2</span><span>3</span><span>4</span><span>5</span></div><span></span></div>")
    A("<div class=\"areas\" role=\"table\" aria-label=\"Highest open level by area\">")
    for a in AREAS:
        lvl = v["area:" + a]
        n_open = sum(1 for r in risks.values() if r.get("area") == a and r.get("level", 0) > 0)
        A("<span role=\"rowheader\">%s</span><div class=\"track\" role=\"cell\"><div class=\"fill\" style=\"width:%d%%\"></div></div><span role=\"cell\">L%d %s · %d open</span>"
          % (e(area_label.get(a, a)), lvl * 20, lvl, LEVEL_WORD[lvl], n_open))
    A("</div></section>")

    # the table in the template's columns
    A("<section class=\"card\" data-kind=\"reference\"><h2>What we believed, tested and now know</h2>")
    A("<p class=\"lead\">One row per belief recorded in %s, in the company template's columns. The same table is in <code>narrative_%s.xlsx</code> and <code>.csv</code>.</p>" % (e(q), e(q)))
    A("<div class=\"scroll\"><table><thead><tr>" + "".join("<th>%s</th>" % e(c) for c in COLUMNS[1:]) + "<th>Record</th></tr></thead><tbody>")
    for row in trows:
        A("<tr>" + "".join("<td>%s</td>" % e(x) for x in row[1:7]) + "<td><code>%s</code><br>%s · %s</td></tr>" % (e(row[7]), e(row[8]), e(row[9])))
    if not trows:
        A("<tr><td colspan=\"7\">No belief was recorded in this quarter.</td></tr>")
    A("</tbody></table></div></section>")

    # the high risks
    A("<section class=\"card\" data-kind=\"reference\"><h2>Open risks at L4 and L5</h2>")
    A("<p class=\"lead\">Each with the test that would close it. A level goes down only when that kind of test is recorded.</p><ul>")
    for r in high:
        A("<li><b>%s · L%d %s</b> — %s <span class=\"xk-claim\" data-claim=\"reference\">proposed</span><br><span class=\"tiny\">Area: %s · owner: %s · closing test: %s</span></li>"
          % (e(r["id"]), r["level"], LEVEL_WORD[r["level"]], e(r["title"]), e(area_label.get(r["area"], r["area"])), e(r["owner"]), e(r.get("closing_test", ""))))
    A("</ul>")
    A("<details data-show=\"read expert\"><summary>Every risk in the register</summary><div class=\"scroll\"><table><thead><tr><th>Risk</th><th>Level</th><th>Area</th><th>What is done</th><th>Closing test</th></tr></thead><tbody>")
    for r in sorted(risks.values(), key=lambda r: r["id"]):
        A("<tr><td>%s · %s</td><td>L%d</td><td>%s</td><td>%s</td><td>%s</td></tr>" % (e(r["id"]), e(r["title"]), r.get("level", 0), e(r["area"]), e(r.get("done", "")), e(r.get("closing_test", ""))))
    A("</tbody></table></div></details></section>")

    # where it breaks, wrong idea
    A("<section class=\"card\" data-kind=\"explanation\"><h2>Where this narrative breaks</h2>")
    A("<div class=\"xk-box breaks\" data-explain=\"breaks\"><ul>")
    A("<li>Every level is a proposal the planning package wrote from SPEC.md §21, until decision D26 confirms the scale and each starting level.</li>")
    A("<li>%d of %d beliefs are untested. Each is a bet, and this page counts it as one; it does not know which bets are right.</li>" % (len(untested), v["beliefs"]))
    A("<li>“Held” means the named test passed, not that the belief is true beyond that test.</li>")
    A("<li>Counts can be moved by splitting or merging risks. The register is reviewed like code.</li></ul></div>")
    A("<div class=\"xk-box wrong\" data-explain=\"wrong-idea\"><div class=\"xk-h\">Common wrong idea</div>“A negative net, more risks opened than closed, is a bad quarter.” "
      "<div data-explain=\"because\">Not by itself. Writing down a bet nobody had named opens a risk, and so does a test that breaks a belief. Both turn something hidden into something known. A bad quarter is one where nothing was tested.</div></div>")
    A("<div class=\"xk-box analogy\" data-explain=\"analogy\"><div class=\"xk-h\">Think of it like</div>A ship's log that records every assumption the navigator makes, and crosses each out only when a landmark confirms or refutes it. "
      "<div data-explain=\"analogy-breaks\"><b>Where the story lies:</b> a navigator's landmark is certain; our tests are not. A held belief can still break under a test nobody has run yet.</div></div>")
    A("</section>")

    # sources
    A("<section class=\"card\" data-kind=\"reference\" data-show=\"read expert\"><h2>Sources</h2><ul>")
    A("<li><b>ECSS-M-ST-80C</b>, Space project management — Risk management (ECSS, 31 July 2008). <span data-explain=\"used-for\">Used for: the idea of scoring a risk by severity and likelihood; our five levels map onto it by decision D26.</span></li>")
    A("<li><b>Quarterly De-risking Narrative</b>, the company's template (<code>derisk/narrative_template.xlsx</code>). <span data-explain=\"used-for\">Used for: the seven columns of the table and the rule that bad news appears here first.</span></li>")
    A("<li><b>SPEC.md §5.13</b>, the ledger's rules. <span data-explain=\"used-for\">Used for: what counts as a move, and when a level may go down.</span></li>")
    A("</ul></section>")

    A("<section class=\"card\" data-kind=\"tutorial\" data-show=\"learn\" data-explain=\"teach-back\"><h2>Explain it back</h2>")
    A("<p>In three lines, as you would to someone joining the team: what does the platform still rest on without proof, which area worries you most, and what single test would help most next quarter?</p><textarea aria-label=\"Your explanation\"></textarea>"
      "<p class=\"tiny\">This stays in your browser. It is for you, not a record.</p></section>")
    A("<p data-explain=\"one-line\" class=\"lead\"><b>In one line:</b> this quarter's ledger, counted; the prose is a person's; every level is a proposal until D26.</p>")
    A("</main>\n</body>\n</html>\n")
    return "\n".join(L)


# ------------------------------------------------------------------ §21 table
T_BEGIN, T_END = "<!-- risks:begin (generated by tools/derisk.py table) -->", "<!-- risks:end -->"


def spec_table(reg):
    L = [T_BEGIN, "| Risk | Area | Level | Why it is real | What is done about it | The test that closes it |", "|---|---|---|---|---|---|"]
    esc = lambda x: str(x).replace("|", "\\|")
    for r in sorted(reg.get("risk", []), key=lambda r: r["id"]):
        L.append("| %s %s | %s | L%d | %s | %s | %s |" % (r["id"], esc(r["title"]), r["area"], r.get("level", 0), esc(r.get("why", "")), esc(r.get("done", "")), esc(r.get("closing_test", ""))))
    L.append(T_END)
    return "\n".join(L)


# ------------------------------------------------------------------ record (the developer team's own changes)
def quarter_now():
    import datetime
    d = datetime.date.today()
    return "Q%d-%02d" % ((d.month - 1) // 3 + 1, d.year % 100)


def record(a, opt, root=None):
    """Writes derisk/beliefs/B-nnn.toml and its moves into derisk/risks.toml, as
    `cargo xtask derisk record` does; nothing is written if the ledger would then
    fail its check."""
    root = root or ROOT
    reg, beliefs, _ = load(root)
    many = lambda name: [a[i + 1] for i, x in enumerate(a) if x == name and i + 1 < len(a)]
    need = {"--area": "area", "--about": "about", "--believed": "believed", "--status": "status",
            "--now-know": "now_know", "--plan-change": "plan_change"}
    missing = [k for k in need if not opt(k)]
    if missing:
        print("record needs %s" % ", ".join(missing))
        return 2
    Q = opt("--quarter") or quarter_now()
    nums = [int(b[2:]) for b in beliefs if re.fullmatch(r"B-\d+", b)]
    bid = "B-%03d" % (max(nums or [0]) + 1)
    cost = opt("--cost-k")
    b = {"schema": "adcs-belief/1", "id": bid, "quarter": Q, "area": opt("--area"),
         "about": [x.strip() for x in opt("--about").split(",") if x.strip()], "request": "", "version": "",
         "recorded_by": "developer team", "believed": opt("--believed"), "status": opt("--status"),
         "tested": opt("--tested") or "", "now_know": opt("--now-know"), "cost_k": float(cost) if cost else "",
         "plan_change": opt("--plan-change"), "previous_issue": opt("--previous-issue") or "",
         "benefit": opt("--benefit") or "", "would_test": opt("--would-test") or "", "risks": [], "move": []}
    risks = {r["id"]: r for r in reg.get("risk", [])}
    nxt = max([int(r[2:]) for r in risks if re.fullmatch(r"R-\d+", r)] or [0]) + 1
    for mv in many("--move"):
        try:
            rid, f, t = mv.split(":")
            f, t = int(f), int(t)
        except ValueError:
            print("--move is R-nn:<from>:<to>, not %r" % mv)
            return 2
        if rid not in risks:
            print("--move names %s, which is not in the register" % rid)
            return 1
        b["move"].append({"risk": rid, "from": f, "to": t})
        r = risks[rid]
        r.setdefault("history", []).append({"quarter": Q, "from": f, "to": t, "belief": bid})
        r["level"] = t
        if t == 0:
            r["closed"] = Q
    for op in many("--open"):
        parts = [x.strip() for x in op.split("|")]
        if len(parts) != 5:
            print("--open is \"<title>|<area>|<level>|<closing test>|<owner>\"")
            return 2
        rid = "R-%02d" % nxt
        nxt += 1
        lvl = int(parts[2])
        reg.setdefault("risk", []).append({"id": rid, "title": parts[0], "area": parts[1], "owner": parts[4],
                                           "why": b["believed"], "done": b["plan_change"], "closing_test": parts[3],
                                           "level": lvl, "opened": Q, "closed": "",
                                           "history": [{"quarter": Q, "from": 0, "to": lvl, "belief": bid}]})
        b["move"].append({"risk": rid, "from": 0, "to": lvl})
    for rid in many("--risk"):   # a risk the belief bears on without moving it
        if rid not in {r["id"] for r in reg.get("risk", [])}:
            print("--risk names %s, which is not in the register" % rid)
            return 1
    b["risks"] = sorted({m["risk"] for m in b["move"]} | set(many("--risk")))
    beliefs[bid] = dict(b, _file=bid + ".toml")
    F = check(reg, beliefs)
    if F:
        for code, where, text in F:
            print("%s %s: %s" % (code, where, text))
        print("not recorded: the ledger would fail its check")
        return 1
    q = lambda v: json.dumps(v, ensure_ascii=False)
    L = []
    for k in ("schema", "id", "quarter", "area", "about", "request", "version", "recorded_by", "believed", "status", "tested",
              "now_know", "cost_k", "plan_change", "previous_issue", "benefit", "would_test", "risks"):
        L.append("%s = %s" % (k, q(b[k]) if b[k] != "" or k != "cost_k" else '""'))
    for m in b["move"]:
        L += ["", "[[move]]", "risk = %s" % q(m["risk"]), "from = %d" % m["from"], "to = %d" % m["to"]]
    open(os.path.join(root, "derisk", "beliefs", bid + ".toml"), "w", encoding="utf-8").write("\n".join(L) + "\n")
    # the register: append the moves as history entries, in place, keeping the file's comments
    rp = os.path.join(root, "derisk", "risks.toml")
    text = open(rp, encoding="utf-8").read()
    for m in b["move"]:
        rid = m["risk"]
        if m["from"] == 0 and rid not in {r["id"] for r in load_toml(rp).get("risk", [])}:
            r = next(x for x in reg["risk"] if x["id"] == rid)
            text = text.rstrip("\n") + "\n\n[[risk]]\n" + "\n".join(
                "%s = %s" % (k, q(r[k]) if not isinstance(r[k], int) else r[k]) for k in
                ("id", "title", "area", "owner", "why", "done", "closing_test", "level", "opened", "closed")) + \
                "\nhistory = [{ quarter = %s, from = 0, to = %d, belief = %s }]\n" % (q(Q), r["level"], q(bid))
            continue
        blk = re.search(r'(?ms)^\[\[risk\]\]\nid = "%s"\n.*?(?=^\[\[risk\]\]|\Z)' % re.escape(rid), text)
        body = blk.group(0)
        new = re.sub(r"(?m)^level = \d+$", "level = %d" % m["to"], body, count=1)
        if m["to"] == 0:
            new = re.sub(r'(?m)^closed = ".*"$', 'closed = %s' % q(Q), new, count=1)
        new = re.sub(r"(?m)^(history = \[.*)\]$", lambda h: h.group(1).rstrip() + ", { quarter = %s, from = %d, to = %d, belief = %s }]"
                     % (q(Q), m["from"], m["to"], q(bid)), new, count=1)
        text = text.replace(body, new)
    open(rp, "w", encoding="utf-8").write(text)
    reg2, bel2, _ = load(root)
    F2 = check(reg2, bel2)
    if F2:
        print("recorded, but the files now fail their check (a bug in record): %s" % F2)
        return 1
    print("recorded %s (%s) with %d move(s)" % (bid, b["status"], len(b["move"])))
    return 0


# ------------------------------------------------------------------ selftest
def selftest():
    reg0, bel0, _ = load()
    F0 = check(reg0, bel0)
    if F0:
        print("the ledger itself fails:", F0)
        return 1

    def mut(fn):
        reg, bel = copy.deepcopy(reg0), copy.deepcopy(bel0)
        fn(reg, bel)
        return reg, bel

    R = lambda reg, rid: next(r for r in reg["risk"] if r["id"] == rid)
    cases = [
        ("L01", "a risk id that is not R-nn", lambda reg, bel: R(reg, "R-12").update({"id": "RISK12"})),
        ("L01", "two risks with one id", lambda reg, bel: R(reg, "R-12").update({"id": "R-11"})),
        ("L02", "an area that is not one of the seven", lambda reg, bel: R(reg, "R-12").update({"area": "finance"})),
        ("L02", "a person as owner", lambda reg, bel: R(reg, "R-12").update({"owner": "A. Person"})),
        ("L02", "a level of 7", lambda reg, bel: (R(reg, "R-12").update({"level": 7}))),
        ("L03", "a level that is not where the history ends", lambda reg, bel: R(reg, "R-12").update({"level": 3})),
        ("L03", "a closed risk with no closing quarter", lambda reg, bel: (R(reg, "R-12")["history"].append({"quarter": "Q4-26", "from": 2, "to": 0, "belief": "B-001"}), R(reg, "R-12").update({"level": 0}))),
        ("L04", "a register move no belief records", lambda reg, bel: R(reg, "R-12")["history"].append({"quarter": "Q4-26", "from": 2, "to": 3, "belief": "B-002"}) or R(reg, "R-12").update({"level": 3})),
        ("L04", "a belief move the register does not record", lambda reg, bel: bel["B-003"].setdefault("move", []).append({"risk": "R-13", "from": 3, "to": 4})),
        ("L05", "a level lowered by an untested belief", lambda reg, bel: (R(reg, "R-15")["history"].append({"quarter": "Q3-26", "from": 3, "to": 2, "belief": "B-008"}),
                                                                           R(reg, "R-15").update({"level": 2}), bel["B-008"]["move"].append({"risk": "R-15", "from": 3, "to": 2}))),
        ("L05", "a level lowered with no belief at all", lambda reg, bel: (R(reg, "R-12")["history"].append({"quarter": "Q4-26", "from": 2, "to": 1, "belief": ""}), R(reg, "R-12").update({"level": 1}))),
        ("L04", "a level raised with no belief", lambda reg, bel: (R(reg, "R-12")["history"].append({"quarter": "Q4-26", "from": 2, "to": 3, "belief": ""}), R(reg, "R-12").update({"level": 3}))),
        ("L06", "a belief with no area", lambda reg, bel: bel["B-006"].update({"area": ""})),
        ("L06", "a belief that says nothing about what we now know", lambda reg, bel: bel["B-006"].update({"now_know": " "})),
        ("L07", "a held belief with no test", lambda reg, bel: bel["B-001"].update({"tested": ""})),
        ("L07", "an untested belief with no test that would settle it", lambda reg, bel: bel["B-006"].update({"would_test": ""})),
        ("L08", "a cost that is not a number", lambda reg, bel: bel["B-006"].update({"cost_k": "a lot"})),
        ("L08", "a quarter written as a date", lambda reg, bel: bel["B-006"].update({"quarter": "2026-09"})),
        ("L09", "a belief naming a risk that does not exist", lambda reg, bel: bel["B-006"].update({"risks": ["R-99"]})),
        ("L09", "an open risk with no closing test", lambda reg, bel: R(reg, "R-12").update({"closing_test": ""})),
    ]
    bad = 0
    for code, what, fn in cases:
        reg, bel = mut(fn)
        got = {c for c, _, _ in check(reg, bel)}
        if code not in got:
            print("NOT REFUSED %s: %s (got %s)" % (code, what, sorted(got)))
            bad += 1
    missing = set(CODES) - {c for c, _, _ in cases}
    if missing:
        print("rules with no deliberate mistake: %s" % sorted(missing))
        bad += 1
    # record: a tested belief lowers a risk and opens another; an untested one may not lower
    tmp = tempfile.mkdtemp()
    try:
        shutil.copytree(os.path.join(ROOT, "derisk"), os.path.join(tmp, "derisk"))
        args = ["record", "--area", "model", "--about", "crates/adcs-sim-core", "--believed", "selftest belief",
                "--status", "held", "--tested", "selftest test", "--now-know", "selftest knowledge", "--plan-change", "none",
                "--quarter", "Q4-26", "--move", "R-12:2:1", "--open", "selftest risk|model|2|a selftest closing test|systems"]
        o = lambda name: args[args.index(name) + 1] if name in args else None
        import contextlib
        with contextlib.redirect_stdout(io.StringIO()):
            rc = record(args, o, tmp)
        r2, b2, _ = load(tmp)
        if rc or check(r2, b2) or next(r for r in r2["risk"] if r["id"] == "R-12")["level"] != 1 or not any(r["id"] == "R-18" for r in r2["risk"]):
            print("record did not write a held belief, its move and its new risk cleanly")
            bad += 1
        args2 = ["record", "--area", "model", "--about", "x", "--believed", "b", "--status", "untested", "--would-test", "t",
                 "--now-know", "k", "--plan-change", "p", "--quarter", "Q4-26", "--move", "R-11:4:3"]
        o2 = lambda name: args2[args2.index(name) + 1] if name in args2 else None
        with contextlib.redirect_stdout(io.StringIO()):
            rc2 = record(args2, o2, tmp)
        if rc2 == 0:
            print("record let an untested belief lower a risk")
            bad += 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    q, rows, concl, v = rollup(reg0, bel0)
    if len(rows) != 18 or set(rows) != {t for t, _, _ in ROWS}:
        print("the rollup does not give the 18 risk rows")
        bad += 1
    tmp = tempfile.mkdtemp()
    try:
        a = narrative_html(reg0, bel0, {}, q)
        b = narrative_html(reg0, bel0, {}, q)
        if a != b:
            print("the narrative is not deterministic")
            bad += 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    if bad:
        return 1
    print("selftest ok: %d deliberate mistakes in the ledger refused by their rules; record writes a tested move and refuses an untested one; rollup gives %d rows; the narrative is deterministic" % (len(cases), len(rows)))
    return 0


# ------------------------------------------------------------------ main
def main():
    a = sys.argv[1:]
    if not a or a[0] in ("-h", "--help"):
        print(__doc__)
        return 2
    cmd = a[0]

    def opt(name, default=None):
        return a[a.index(name) + 1] if name in a and a.index(name) + 1 < len(a) else default

    if cmd == "codes":
        for k, t in CODES.items():
            print("%s  %s" % (k, t))
        return 0
    if cmd == "selftest":
        return selftest()
    if cmd == "template":
        out = opt("--out")
        if not out:
            print("usage: tools/derisk.py template --out <file.xlsx>")
            return 2
        write_xlsx(out, [], blank=True)
        print("wrote %s" % out)
        return 0
    reg, beliefs, prose = load()
    if cmd == "check":
        F = check(reg, beliefs)
        for code, where, text in F:
            print("%s %s: %s" % (code, where, text))
        print("%d finding(s) · %d risks · %d beliefs" % (len(F), len(reg.get("risk", [])), len(beliefs)))
        return 1 if F else 0
    if cmd == "rollup":
        q, rows, concl, v = rollup(reg, beliefs, opt("--quarter"))
        if "--json" in a:
            print(json.dumps({"quarter": q, "rows": rows, "conclusion": concl}, indent=1, sort_keys=True))
            return 0
        print("Quarter %s — the values the risk rows take (supplier derisk):" % q)
        for tid, label, _ in ROWS:
            print("  %-6s %-40s %s" % (tid, label, rows[tid]))
        print("The conclusion rows, which the engine computes from those (shown here for reading, not supplied):")
        for tid, label, how in CONCLUSION:
            val = concl[tid]
            print("  %-6s %-40s %s   (%s)" % (tid, label, "Undefined" if val is None else (round(val, 4) if isinstance(val, float) else val), how))
        return 0
    if cmd == "narrative":
        q, out = opt("--quarter") or latest_quarter(reg, beliefs), opt("--out")
        if not out:
            print("usage: tools/derisk.py narrative --quarter Q3-26 --out <dir>")
            return 2
        if not QUARTER.match(q):
            print("quarter %r is not Qn-YY" % q)
            return 2
        os.makedirs(out, exist_ok=True)
        page = narrative_html(reg, beliefs, prose, q)
        open(os.path.join(out, "narrative_%s.html" % q), "w", encoding="utf-8").write(page)
        rows = table_rows(beliefs, q, {r["id"]: r for r in reg.get("risk", [])})
        write_csv(os.path.join(out, "narrative_%s.csv" % q), rows)
        try:
            write_xlsx(os.path.join(out, "narrative_%s.xlsx" % q), rows)
            xl = ", .xlsx"
        except ImportError:
            xl = " (no .xlsx: openpyxl is not installed)"
        print("wrote narrative_%s.html, .csv%s in %s (%d beliefs)" % (q, xl, out, len(rows)))
        return 0
    if cmd == "table":
        p = opt("--file")
        if not p:
            p = os.path.join(ROOT, "spec", "19_phases.md")
            if not os.path.exists(p):
                p = os.path.join(ROOT, "docs", "RISKS.md")
        if not os.path.exists(p) and "--check" not in a:
            os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
            open(p, "w", encoding="utf-8").write("# Risks carried\n\nGenerated from `derisk/risks.toml` by `cargo xtask derisk table`; "
                                                "CI fails when it differs (SPEC.md §5.13, §21).\n\n%s\n%s\n" % (T_BEGIN, T_END))
        text = open(p, encoding="utf-8").read() if os.path.exists(p) else ""
        if T_BEGIN not in text:
            print("%s has no %s marker" % (p, T_BEGIN))
            return 1
        new = re.sub(re.escape(T_BEGIN) + r".*?" + re.escape(T_END), lambda m: spec_table(reg), text, count=1, flags=re.S)
        if "--check" in a:
            if new != text:
                print("STALE: the risk table in %s differs from derisk/risks.toml — run python3 tools/derisk.py table" % os.path.relpath(p, ROOT))
                return 1
            print("the risk table in %s is current" % os.path.relpath(p, ROOT))
            return 0
        open(p, "w", encoding="utf-8").write(new)
        print("wrote the risk table into %s" % os.path.relpath(p, ROOT))
        return 0
    if cmd == "record":
        return record(a, opt)
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main())
